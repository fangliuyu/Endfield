
import logging
import traceback
from collections import OrderedDict
from pathlib import Path

import pandas as pd
from PySide6.QtCore import QObject, Signal
from openpyxl.reader.excel import load_workbook
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter

from lib.deviceCheck import check_error


class build_summary_work(QObject):
    finish_signal = Signal(str, bool, str)  # type, bool, msg

    def __init__(self, parent, data_folder: Path, logger: logging.Logger):
        super().__init__(parent)
        self.logger = logger
        self.data_folder = data_folder

    @check_error
    def run(self):
        self.combine_csv_results()

    def combine_csv_results(self):
        try:
            self.logger.debug(f"combine folder: {self.data_folder.as_posix()}")
            folder_list = [entry for entry in self.data_folder.iterdir() if entry.is_dir()]
            self.logger.debug(f"combine file number: {len(folder_list)}")
            dataframes = OrderedDict()
            folder_list.sort(key=lambda x: x.Name)
            for folder in folder_list:
                for folder_file in folder.iterdir():
                    if "Measurement Data.xlsx" not in folder_file.name:
                        continue
                    try:
                        test_name = folder.name
                        self.logger.debug(f"combine file: {test_name}")
                        test_name = test_name.replace(f"{self.data_folder.name}-", "")
                        self.logger.debug(f"after replace combine file: {test_name}")
                        df = pd.read_excel(folder_file, engine='openpyxl')
                        columns = list(set(df.columns))
                        if "Link" in columns:
                            df.drop("Link", axis=1, inplace=True)
                        if "Result" in columns:
                            df.drop("Result", axis=1, inplace=True)
                        df = df.rename(columns={
                            'Measured Value': test_name
                        })

                        # 将 DataFrames 添加到列表中
                        dataframes[test_name] = df
                    except Exception as e:
                        self.logger.error(f"读取文件 {folder.name} 时发生错误：{str(e)}")
                        continue
            aligned_df = pd.DataFrame()
            front_columns = ["Net", "Measurement", "Lower Limit", "Lower Warning", "Upper Warning", "Upper Limit", "Unit"]

            # First find a DataFrame that has all front_columns to initialize aligned_df
            initialized = False
            for name, df in dataframes.items():
                self.logger.debug(name)
                if all(col in df.columns for col in front_columns):
                    aligned_df = df.copy()
                    initialized = True
                    break

            if not initialized:
                raise ValueError("None of the input DataFrames contain all required front columns")

            # Now merge remaining DataFrames
            for name, df in dataframes.items():
                if df.equals(aligned_df):  # Skip the one we used for initialization
                    continue
                if not all(col in df.columns for col in front_columns):
                    print(f"Warning: DataFrame {name} missing some front columns, skipping merge")
                    continue
                aligned_df = pd.merge(aligned_df, df, how="outer", on=front_columns)

            values_columns = [col for col in dataframes.keys() if col in aligned_df.columns]
            file_path = self.data_folder / "Summary Results.xlsx"

            # Ensure all columns exist before exporting
            missing_cols = [col for col in front_columns + values_columns if col not in aligned_df.columns]
            if missing_cols:
                raise ValueError(f"Missing columns in final DataFrame: {missing_cols}")

            # aligned_df[front_columns + values_columns].sort_index().to_excel(file_path, index=False)
            aligned_df[front_columns + values_columns].to_excel(file_path, index=False)

            wb = load_workbook(file_path)
            ws = wb.active
            ws.title = "Measurement Summary"
            green_fill = PatternFill(
                start_color="82E0AA",
                end_color="82E0AA",
                fill_type="solid"
            )
            red_fill = PatternFill(
                start_color="F1948A",
                end_color="F1948A",
                fill_type="solid"
            )
            yellow_fill = PatternFill(
                start_color="F9E79F",
                end_color="F9E79F",
                fill_type="solid"
            )
            columns_list = []
            for columns in ws[1]:
                columns_list.append(columns.value)
            columns_dict = {}
            test_loc_list = []
            for i in range(len(columns_list)):
                columns_dict[columns_list[i]] = i
                if columns_list[i] not in front_columns:
                    test_loc_list.append(i)
            pass_loc = []
            fail_loc = []
            warn_loc = []
            for index in range(2, ws.max_row + 1):
                row_data = []
                for row in ws[index]:
                    row_data.append(row.value)
                for test_loc_x in test_loc_list:
                    self.logger.debug(
                        [
                            row_data[columns_dict["Lower Limit"]],
                            row_data[columns_dict["Lower Warning"]],
                            row_data[test_loc_x],
                            row_data[columns_dict["Upper Warning"]],
                            row_data[columns_dict["Upper Limit"]]
                        ]
                    )
                    if row_data[test_loc_x] == "/" or row_data[test_loc_x] is None:
                        self.logger.debug(f"choose to /")
                        continue
                    if row_data[columns_dict["Lower Limit"]] != "-":
                        if float(row_data[test_loc_x]) < float(row_data[columns_dict["Lower Limit"]]):
                            fail_loc.append([index, test_loc_x])
                            self.logger.debug(f"choose to fail to lower")
                            continue
                    if row_data[columns_dict["Upper Limit"]] != "-":
                        if float(row_data[test_loc_x]) > float(row_data[columns_dict["Upper Limit"]]):
                            fail_loc.append([index, test_loc_x])
                            self.logger.debug(f"choose to fail to upper")
                            continue
                    if row_data[columns_dict["Upper Warning"]] == "-":
                        if row_data[columns_dict["Lower Warning"]] == "-":
                            pass_loc.append([index, test_loc_x])
                            self.logger.debug(f"choose to pass to no spc")
                        else:
                            # Upper Warning 为 '-'，只需检查是否大于 Lower Warning
                            if float(row_data[test_loc_x]) > float(row_data[columns_dict["Lower Warning"]]):
                                pass_loc.append([index, test_loc_x])
                                self.logger.debug(f"choose to pass (only lower check)")
                            else:
                                warn_loc.append([index, test_loc_x])
                                self.logger.debug(f"choose to marginal (lower violation)")
                    elif row_data[columns_dict["Lower Warning"]] == "-":
                        # Lower Warning 为 '-'，只需检查是否小于 Upper Warning
                        if float(row_data[test_loc_x]) < float(row_data[columns_dict["Upper Warning"]]):
                            pass_loc.append([index, test_loc_x])
                            self.logger.debug(f"choose to pass (only upper check)")
                        else:
                            warn_loc.append([index, test_loc_x])
                            self.logger.debug(f"choose to marginal (upper violation)")
                    else:
                        # 正常情况：检查是否在 Lower Warning 和 Upper Warning 之间
                        if float(row_data[columns_dict["Lower Warning"]]) < float(row_data[test_loc_x]) < float(row_data[columns_dict["Upper Warning"]]):
                            pass_loc.append([index, test_loc_x])
                            self.logger.debug(f"choose to pass")
                        else:
                            warn_loc.append([index, test_loc_x])
                            self.logger.debug(f"choose to marginal")

            for loc in pass_loc:
                ws.cell(row=loc[0], column=loc[1] + 1).fill = green_fill
            for loc in fail_loc:
                ws.cell(row=loc[0], column=loc[1] + 1).fill = red_fill
            for loc in warn_loc:
                ws.cell(row=loc[0], column=loc[1] + 1).fill = yellow_fill

            ws.auto_filter.ref = ws.dimensions
            ws.column_dimensions['A'].width = 20  # "Net."
            ws.column_dimensions['B'].width = 30  # "Measurement"
            ws.column_dimensions['C'].width = 20  # "Lower Limit"
            ws.column_dimensions['D'].width = 20  # "Lower Warning"
            ws.column_dimensions['E'].width = 20  # "Upper Warning"
            ws.column_dimensions['F'].width = 20  # "Upper Limit"
            ws.column_dimensions['G'].width = 15  # "Unit"
            for i in range(8, ws.max_column + 1):
                ws.column_dimensions[get_column_letter(i)].width = 30
            wb.save(file_path)
            self.logger.info(f"combine report excel files successfully")

        except Exception as e:
            self.logger.debug(traceback.format_exc())
            self.logger.error(f"整个程序在处理 Excel 文件时发生错误：{str(e)}")


def draw_pass_color(row):
    if row["Result"] == "Pass":
        return ["background-color: green"] * len(row)
