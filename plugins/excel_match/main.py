
import logging
import os
import subprocess
import sys
import traceback
from os import PathLike
from pathlib import Path
from typing import Optional, Union

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import PatternFill

from PySide6.QtWidgets import (
    QApplication, QMainWindow, QWidget, QVBoxLayout, QHBoxLayout,
    QPushButton, QLabel, QFrame, QComboBox, QFileDialog,
    QPlainTextEdit, QTabWidget, QMessageBox, QTableWidget,
    QTableWidgetItem, QHeaderView, QSpinBox, QLineEdit,
    QDialog, QListWidget, QDialogButtonBox,
)

from lib.customlog import create_logger


# =============================================================================
# Helper: unlock password-protected Excel
# =============================================================================

def unlock_excel(filepath: str, password: str, temp_dir: Path) -> str:
    import msoffcrypto
    file_name = os.path.basename(filepath)
    with open(filepath, 'rb') as f:
        excel_file = msoffcrypto.OfficeFile(f)
        excel_file.load_key(password)
        out = temp_dir / file_name
        with open(out, "wb") as fout:
            excel_file.decrypt(fout)
    return str(out)


def open_excel_with_fallback(filepath: str, temp_dir: Path, parent=None) -> tuple[pd.ExcelFile, str]:
    """尝试打开Excel，遇到加密则弹窗询问密码，返回 (ExcelFile, 实际路径)"""
    try:
        ef = pd.ExcelFile(filepath, engine='openpyxl')
        return ef, filepath
    except Exception:
        pass

    from PySide6.QtWidgets import QInputDialog
    file_name = os.path.basename(filepath)
    password, ok = QInputDialog.getText(parent, f"文件需要密码", f"{file_name}\n请输入密码:",
                                        echo=QLineEdit.EchoMode.Password)
    if not ok or not password:
        raise PermissionError("用户取消密码输入")

    real_path = unlock_excel(filepath, password, temp_dir)
    return pd.ExcelFile(real_path, engine='openpyxl'), real_path


# =============================================================================
# Helper: check strikethrough
# =============================================================================

def has_strikethrough(cell) -> bool:
    try:
        return cell.font.strikethrough
    except AttributeError:
        return False


def read_excel_sheet(file_path: str, sheet_name: str) -> pd.DataFrame:
    """读取Excel时跳过带删除线的行"""
    wb = load_workbook(file_path, data_only=True)
    ws = wb[sheet_name]
    headers = [cell.value for cell in ws[1]]
    rows = []
    for row in ws.iter_rows(min_row=2, values_only=False):
        if not any(has_strikethrough(cell) for cell in row):
            rows.append([cell.value for cell in row])
    return pd.DataFrame(rows, columns=headers)


def consolidate_sheets(sheets_data: dict[str, pd.DataFrame],
                       sn_column: str, quantity_column: str,
                       data_row: int) -> pd.DataFrame:
    """合并多个sheet的数据"""
    consolidated_df = pd.DataFrame()
    total_quantity = {}

    for sheet_name, df in sheets_data.items():
        if consolidated_df.empty:
            consolidated_df = pd.DataFrame(columns=df.columns)
            for index, row in df.iterrows():
                if int(index) < data_row:
                    continue
                consolidated_df.loc[len(consolidated_df)] = row.copy()

        for index, row in df.iterrows():
            if int(index) < data_row:
                continue
            sn = row[sn_column]
            if pd.isna(sn):
                continue
            quantity = pd.to_numeric(row[quantity_column], errors='coerce')
            if pd.isna(quantity):
                quantity = 0
            quantity = float(quantity)

            if sn in consolidated_df[sn_column].values:
                idx = consolidated_df[consolidated_df[sn_column] == sn].index[0]
                total_quantity[sn] = total_quantity.get(sn, 0) + quantity
                consolidated_df.loc[idx, f"Quantity in\n{sheet_name}"] = quantity
                consolidated_df.loc[idx, quantity_column] = total_quantity[sn]
            else:
                new_row = row.copy()
                new_row[f"Quantity in\n{sheet_name}"] = quantity
                new_row[quantity_column] = quantity
                total_quantity[sn] = quantity
                consolidated_df.loc[len(consolidated_df)] = new_row

    consolidated_df.rename(columns={quantity_column: f'total-{quantity_column}'}, inplace=True)
    return consolidated_df


# =============================================================================
# Dialog: Sheet selection popup
# =============================================================================

class SheetSelectDialog(QDialog):
    def __init__(self, filepath: str, sheet_list: list[str], parent=None):
        super().__init__(parent)
        self.selected_sheet = sheet_list[0] if sheet_list else ""
        self.setWindowTitle("选择工作表")
        self.setModal(True)
        self.resize(400, 200)

        layout = QVBoxLayout(self)
        layout.addWidget(QLabel(f"文件: {os.path.basename(filepath)}"))
        layout.addWidget(QLabel("选择工作表:"))

        self.list_widget = QListWidget()
        self.list_widget.addItems(sheet_list)
        if sheet_list:
            self.list_widget.setCurrentRow(0)
        layout.addWidget(self.list_widget)

        btn_box = QDialogButtonBox(QDialogButtonBox.StandardButton.Ok |
                                   QDialogButtonBox.StandardButton.Cancel)
        btn_box.accepted.connect(self.accept)
        btn_box.rejected.connect(self.reject)
        layout.addWidget(btn_box)

    def accept(self):
        item = self.list_widget.currentItem()
        if item:
            self.selected_sheet = item.text()
        super().accept()


# =============================================================================
# Tab 1: Excel 整合 (原 integrate.py)
# =============================================================================

class IntegrateTab(QWidget):
    def __init__(self, temp_dir: Path, output_dir: Path, parent=None):
        super().__init__(parent)
        self.temp_dir = temp_dir
        self.output_dir = output_dir
        self.file_records: list[tuple[str, str]] = []  # (filepath, sheet_name)

        layout = QVBoxLayout(self)
        layout.setSpacing(5)

        # Instruction
        instr = QLabel(
            "点击\"+\"添加Excel文件，支持同时添加多个。点击\"-\"从下至上逐个删除。\n"
            "\"索引标签\"指自动递增的列标题。\"匹配标签\"指用于作整合对象的列标题。\n"
            "\"数量标签\"指整合数量的列标题。\"忽略标签\"指不会导出的列标题。\n"
            "\"Output Name\"填写输出文件名，文件会保存在result文件夹中。\n"
            "注意：输出将以每个sheet名字显示每个Excel的各自数量。"
        )
        instr.setStyleSheet("background-color: #555; color: #fff; padding: 4px;")
        layout.addWidget(instr)

        # File list table
        self.file_table = QTableWidget(0, 2)
        self.file_table.setHorizontalHeaderLabels(["File", "Sheet"])
        self.file_table.horizontalHeader().setSectionResizeMode(0, QHeaderView.ResizeMode.Stretch)
        self.file_table.horizontalHeader().setSectionResizeMode(1, QHeaderView.ResizeMode.ResizeToContents)
        self.file_table.verticalHeader().setVisible(False)
        layout.addWidget(self.file_table, 1)

        # File buttons
        btn_row = QHBoxLayout()
        layout.addLayout(btn_row)
        self.btn_add = QPushButton("+ 添加文件")
        self.btn_add.clicked.connect(self.add_files)
        btn_row.addWidget(self.btn_add)
        self.btn_remove = QPushButton("- 删除最后")
        self.btn_remove.clicked.connect(self.remove_last)
        btn_row.addWidget(self.btn_remove)
        self.btn_clear = QPushButton("清空")
        self.btn_clear.clicked.connect(self.clear_all)
        btn_row.addWidget(self.btn_clear)

        # Separator
        layout.addWidget(self._sep())

        # Row config
        row_cfg = QHBoxLayout()
        layout.addLayout(row_cfg)
        row_cfg.addWidget(QLabel("列名行数:"))
        self.title_row_spin = QSpinBox()
        self.title_row_spin.setRange(1, 50)
        self.title_row_spin.setValue(1)
        row_cfg.addWidget(self.title_row_spin)
        row_cfg.addWidget(QLabel("数据起始行:"))
        self.data_row_spin = QSpinBox()
        self.data_row_spin.setRange(1, 50)
        self.data_row_spin.setValue(2)
        row_cfg.addWidget(self.data_row_spin)
        self.btn_refresh = QPushButton("刷新列名")
        self.btn_refresh.clicked.connect(self.refresh_columns)
        row_cfg.addWidget(self.btn_refresh)

        # Column combos
        combo_row = QHBoxLayout()
        layout.addLayout(combo_row)
        combo_row.addWidget(QLabel("索引标签:"))
        self.index_combo = QComboBox()
        combo_row.addWidget(self.index_combo, 1)
        combo_row.addWidget(QLabel("匹配标签:"))
        self.match_combo = QComboBox()
        combo_row.addWidget(self.match_combo, 1)
        combo_row.addWidget(QLabel("数量标签:"))
        self.variable_combo = QComboBox()
        combo_row.addWidget(self.variable_combo, 1)
        combo_row.addWidget(QLabel("忽略标签:"))
        self.neglect_combo = QComboBox()
        combo_row.addWidget(self.neglect_combo, 1)

        # Separator
        layout.addWidget(self._sep())

        # Output + Start
        exec_row = QHBoxLayout()
        layout.addLayout(exec_row)
        exec_row.addWidget(QLabel("Output Name:"))
        self.output_name = QLineEdit("output")
        exec_row.addWidget(self.output_name, 1)
        self.btn_start = QPushButton("开始整合")
        self.btn_start.clicked.connect(self.run)
        self.btn_start.setMinimumHeight(32)
        exec_row.addWidget(self.btn_start)

        # Log
        self.log_output = QPlainTextEdit()
        self.log_output.setReadOnly(True)
        self.log_output.setMaximumBlockCount(200)
        self.log_output.setStyleSheet("font-family: 'Courier New', monospace; font-size: 12px;")
        layout.addWidget(self.log_output, 1)

    @staticmethod
    def _sep():
        s = QFrame()
        s.setFrameShape(QFrame.Shape.HLine)
        s.setFrameShadow(QFrame.Shadow.Sunken)
        return s

    def log(self, msg: str):
        self.log_output.appendPlainText(msg)
        QApplication.processEvents()

    def add_files(self):
        paths, _ = QFileDialog.getOpenFileNames(self, "选择Excel文件", "",
                                                "Excel文件 (*.xlsx *.xls);;所有文件 (*)")
        if not paths:
            return

        for fp in paths:
            try:
                ef, actual = open_excel_with_fallback(fp, self.temp_dir, self)
            except Exception as e:
                self.log(f"跳过 {os.path.basename(fp)}: {e}")
                continue

            sheets = ef.sheet_names
            if len(sheets) == 1:
                self.file_records.append((actual, sheets[0]))
            else:
                dlg = SheetSelectDialog(actual, sheets, self)
                if dlg.exec() == QDialog.DialogCode.Accepted:
                    self.file_records.append((actual, dlg.selected_sheet))
                else:
                    continue

        self._update_table()

    def _update_table(self):
        self.file_table.setRowCount(len(self.file_records))
        for i, (fp, sheet) in enumerate(self.file_records):
            self.file_table.setItem(i, 0, QTableWidgetItem(os.path.basename(fp)))
            self.file_table.setItem(i, 1, QTableWidgetItem(sheet))
        if self.file_records:
            self.refresh_columns()

    def remove_last(self):
        if self.file_records:
            self.file_records.pop()
            self._update_table()

    def clear_all(self):
        self.file_records.clear()
        self.file_table.setRowCount(0)

    def refresh_columns(self):
        if not self.file_records:
            return
        fp, sheet = self.file_records[0]
        try:
            title_handle = self.title_row_spin.value() - 1
            df = pd.read_excel(fp, sheet_name=sheet, header=title_handle).dropna(how='all')
        except Exception:
            self.log(f"读取列名失败: {traceback.format_exc()}")
            QMessageBox.warning(self, "错误", "无法读取文件列名，可能是加密文件或格式错误。")
            return

        titles = [""] + df.columns.tolist()
        for combo in [self.index_combo, self.match_combo, self.variable_combo, self.neglect_combo]:
            combo.clear()
            combo.addItems(titles)
        if len(titles) > 1:
            self.index_combo.setCurrentIndex(1)
            if len(titles) > 2:
                self.match_combo.setCurrentIndex(2)
                if len(titles) > 3:
                    self.variable_combo.setCurrentIndex(3)

    def run(self):
        if not self.file_records:
            QMessageBox.warning(self, "提示", "请先添加要整合的Excel文件")
            return

        self.log_output.clear()
        self.log("开始整合...")
        self.btn_start.setEnabled(False)

        try:
            self._do_run()
        except Exception:
            self.log(traceback.format_exc())
            QMessageBox.critical(self, "错误", f"整合失败: {traceback.format_exc()}")
        finally:
            self.btn_start.setEnabled(True)

    def _do_run(self):
        sheets_data = {}
        for fp, sheet in self.file_records:
            try:
                title_handle = self.title_row_spin.value() - 1
                df = pd.read_excel(fp, sheet_name=sheet, header=title_handle)
                df.dropna(subset=[self.match_combo.currentText()], inplace=True)
            except Exception:
                self.log(f"跳过 {os.path.basename(fp)}: {traceback.format_exc()}")
                QMessageBox.warning(self, "警告", f"{os.path.basename(fp)} 打开失败，已跳过。")
                continue

            key = sheet if sheet not in sheets_data else f"{sheet}-{len(sheets_data)}"
            sheets_data[key] = df
            self.log(f"已加载: {os.path.basename(fp)} / {sheet}")

        if not sheets_data:
            self.log("没有可处理的数据")
            return

        match_name = self.match_combo.currentText()
        variable_name = self.variable_combo.currentText()
        ignore_name = self.neglect_combo.currentText()
        data_offset = self.data_row_spin.value() - self.title_row_spin.value()

        result_df = consolidate_sheets(sheets_data, match_name, variable_name, data_offset)

        try:
            if ignore_name:
                result_df.drop(columns=ignore_name, inplace=True)
        except Exception:
            pass

        result_df.dropna(subset=[match_name], inplace=True)
        result_df[match_name] = result_df[match_name].astype(str)
        result_df[self.index_combo.currentText()] = range(1, len(result_df) + 1)

        out_path = self.output_dir / (self.output_name.text() + ".xlsx")
        with pd.ExcelWriter(out_path, engine='openpyxl') as writer:
            result_df.to_excel(writer, sheet_name="Integrate", index=False)

        self.log(f"输出文件: {out_path}")
        QMessageBox.information(self, "完成", f"整合完成!\n{out_path}")
        subprocess.run(['open', str(self.output_dir)])


# =============================================================================
# Tab 2: Excel 对比 (原 match.py)
# =============================================================================

COLOR_MAP = {
    'red': PatternFill(start_color='FFFF0000', end_color='FFFF0000', fill_type='solid'),
    'yellow': PatternFill(start_color='FFFFFF00', end_color='FFFFFF00', fill_type='solid'),
    'green': PatternFill(start_color='FF00FF00', end_color='FF00FF00', fill_type='solid'),
}


class MatchTab(QWidget):
    def __init__(self, temp_dir: Path, output_dir: Path, parent=None):
        super().__init__(parent)
        self.temp_dir = temp_dir
        self.output_dir = output_dir
        self.base_path = ""
        self.match_path = ""
        self.base_sheet = ""
        self.match_sheet = ""

        layout = QVBoxLayout(self)
        layout.setSpacing(5)

        # Instruction
        instr = QLabel(
            "\"Base Excel\"指作为基准的文件,\"Match Excel\"指要比对的文件。\n"
            "删除的行为红色，改动的为黄色，新增的为绿色。\n"
            "\"Output Name\"填写输出文件名，文件会保存在\"result\"文件夹中。"
        )
        instr.setStyleSheet("background-color: #555; color: #fff; padding: 4px;")
        layout.addWidget(instr)

        # Base file
        base_frame = QFrame()
        base_frame.setFrameStyle(QFrame.Shape.Box)
        layout.addWidget(base_frame)
        base_layout = QHBoxLayout(base_frame)
        base_layout.addWidget(QLabel("Base Excel:"))
        self.base_label = QLabel("未选择")
        self.base_label.setStyleSheet("color: #888;")
        base_layout.addWidget(self.base_label, 1)
        self.btn_base = QPushButton("选择")
        self.btn_base.clicked.connect(self.choose_base)
        base_layout.addWidget(self.btn_base)
        base_layout.addWidget(QLabel("Sheet:"))
        self.base_sheet_label = QLabel("-")
        base_layout.addWidget(self.base_sheet_label)

        # Match file
        match_frame = QFrame()
        match_frame.setFrameStyle(QFrame.Shape.Box)
        layout.addWidget(match_frame)
        match_layout = QHBoxLayout(match_frame)
        match_layout.addWidget(QLabel("Match Excel:"))
        self.match_label = QLabel("未选择")
        self.match_label.setStyleSheet("color: #888;")
        match_layout.addWidget(self.match_label, 1)
        self.btn_match = QPushButton("选择")
        self.btn_match.clicked.connect(self.choose_match)
        match_layout.addWidget(self.btn_match)
        match_layout.addWidget(QLabel("Sheet:"))
        self.match_sheet_label = QLabel("-")
        match_layout.addWidget(self.match_sheet_label)

        layout.addWidget(self._sep())

        # Column combos
        combo_row = QHBoxLayout()
        layout.addLayout(combo_row)
        combo_row.addWidget(QLabel("索引标签:"))
        self.index_combo = QComboBox()
        combo_row.addWidget(self.index_combo, 1)
        combo_row.addWidget(QLabel("匹配标签:"))
        self.match_combo = QComboBox()
        combo_row.addWidget(self.match_combo, 1)
        combo_row.addWidget(QLabel("对比标签:"))
        self.variable_combo = QComboBox()
        combo_row.addWidget(self.variable_combo, 1)

        layout.addWidget(self._sep())

        # Output + Start
        exec_row = QHBoxLayout()
        layout.addLayout(exec_row)
        exec_row.addWidget(QLabel("Output Name:"))
        self.output_name = QLineEdit("output")
        exec_row.addWidget(self.output_name, 1)
        self.btn_start = QPushButton("开始对比")
        self.btn_start.clicked.connect(self.run)
        self.btn_start.setMinimumHeight(32)
        exec_row.addWidget(self.btn_start)

        # Log
        self.log_output = QPlainTextEdit()
        self.log_output.setReadOnly(True)
        self.log_output.setMaximumBlockCount(200)
        self.log_output.setStyleSheet("font-family: 'Courier New', monospace; font-size: 12px;")
        layout.addWidget(self.log_output, 1)

    @staticmethod
    def _sep():
        s = QFrame()
        s.setFrameShape(QFrame.Shape.HLine)
        s.setFrameShadow(QFrame.Shadow.Sunken)
        return s

    def log(self, msg: str):
        self.log_output.appendPlainText(msg)
        QApplication.processEvents()

    @staticmethod
    def _load_file_columns(filepath: str, sheet: str):
        """读取列名填充combo"""
        try:
            df = pd.read_excel(filepath, sheet_name=sheet, engine='openpyxl').dropna(how='all')
        except Exception:
            return []
        titles = df.columns.tolist()
        return titles

    def _choose_file(self, target: str):
        fp, _ = QFileDialog.getOpenFileName(self, f"选择{target}Excel文件", "",
                                            "Excel文件 (*.xlsx *.xls);;所有文件 (*)")
        if not fp:
            return None

        try:
            ef, actual = open_excel_with_fallback(fp, self.temp_dir, self)
        except Exception as e:
            QMessageBox.warning(self, "错误", f"打开失败: {e}")
            return None

        sheets = ef.sheet_names
        if len(sheets) == 1:
            return actual, sheets[0]
        else:
            dlg = SheetSelectDialog(actual, sheets, self)
            if dlg.exec() == QDialog.DialogCode.Accepted:
                return actual, dlg.selected_sheet
        return None

    def choose_base(self):
        result = self._choose_file("基准")
        if result:
            self.base_path, self.base_sheet = result
            self.base_label.setText(os.path.basename(self.base_path))
            self.base_label.setStyleSheet("color: #000;")
            self.base_sheet_label.setText(self.base_sheet)
            self._update_combos()

    def choose_match(self):
        result = self._choose_file("比对")
        if result:
            self.match_path, self.match_sheet = result
            self.match_label.setText(os.path.basename(self.match_path))
            self.match_label.setStyleSheet("color: #000;")
            self.match_sheet_label.setText(self.match_sheet)
            self._update_combos()

    def _update_combos(self):
        fp = self.base_path or self.match_path
        if not fp:
            return
        sheet = self.base_sheet if self.base_path else self.match_sheet
        titles = self._load_file_columns(fp, sheet)
        if not titles:
            return
        for combo in [self.index_combo, self.match_combo, self.variable_combo]:
            current = combo.currentText()
            combo.clear()
            combo.addItems(titles)
            if current:
                combo.setCurrentText(current)
        if not self.index_combo.currentText() and len(titles) > 0:
            self.index_combo.setCurrentIndex(0)
        if not self.match_combo.currentText() and len(titles) > 1:
            self.match_combo.setCurrentIndex(1)
        if not self.variable_combo.currentText() and len(titles) > 3:
            self.variable_combo.setCurrentIndex(3)

    def run(self):
        if not self.base_path or not self.match_path:
            QMessageBox.warning(self, "提示", "请先选择基准文件和比对文件")
            return

        self.log_output.clear()
        self.log("开始对比...")
        self.btn_start.setEnabled(False)

        try:
            self._do_run()
        except Exception:
            self.log(traceback.format_exc())
            QMessageBox.critical(self, "错误", f"对比失败: {traceback.format_exc()}")
        finally:
            self.btn_start.setEnabled(True)

    def _do_run(self):
        df1 = pd.read_excel(self.base_path, sheet_name=self.base_sheet, engine='openpyxl')
        df2 = pd.read_excel(self.match_path, sheet_name=self.match_sheet, engine='openpyxl')

        gpn = self.match_combo.currentText()
        qty = self.variable_combo.currentText()
        index_col = self.index_combo.currentText()

        self.log(f"基准: {os.path.basename(self.base_path)} / {self.base_sheet}")
        self.log(f"比对: {os.path.basename(self.match_path)} / {self.match_sheet}")
        self.log(f"匹配列: {gpn}, 对比列: {qty}")

        comparison_df = df1.copy()
        comparison_df.dropna(subset=[gpn], inplace=True)
        comparison_df['Background-Color'] = ''

        for _, row in df2.iterrows():
            sn = row[gpn]
            if pd.isna(sn):
                continue
            quantity = row[qty]

            if sn in df1[gpn].values:
                idx = df1[df1[gpn] == sn].index[0]
                std_qty = df1.loc[idx, qty]
                if std_qty != quantity:
                    comparison_df.loc[idx, 'Background-Color'] = 'yellow'
                    comparison_df.loc[idx, qty] = quantity
            else:
                new_row = row.copy()
                new_row['Background-Color'] = 'green'
                comparison_df.loc[len(comparison_df)] = new_row

        deleted_sns = set(df1[gpn]) - set(df2[gpn])
        for sn in deleted_sns:
            idx = df1[df1[gpn] == sn].index[0]
            comparison_df.loc[idx, 'Background-Color'] = 'red'

        comparison_df.dropna(subset=[gpn], inplace=True)
        comparison_df[gpn] = comparison_df[gpn].astype(str)
        comparison_df[index_col] = range(1, len(comparison_df) + 1)

        # Write colored output
        wb = load_workbook(self.base_path)
        ws = wb[self.base_sheet]

        header_offset = 1
        start_row = header_offset + 1
        headers = [col for col in comparison_df.columns if col != 'Background-Color']
        for col_num, col_name in enumerate(headers, start=1):
            ws.cell(row=header_offset, column=col_num, value=col_name)

        for idx, row in comparison_df.iterrows():
            for col in comparison_df.columns:
                if col == 'Background-Color':
                    continue
                cell = ws.cell(row=start_row + int(idx),
                               column=comparison_df.columns.get_loc(col) + 1)
                cell.value = row[col]
                color = row.get('Background-Color', '')
                if color in COLOR_MAP:
                    cell.fill = COLOR_MAP[color]

        out_path = self.output_dir / (self.output_name.text() + ".xlsx")
        wb.save(str(out_path))
        self.log(f"输出文件: {out_path}")
        QMessageBox.information(self, "完成", f"对比完成!\n{out_path}")
        subprocess.run(['open', str(self.output_dir)])


# =============================================================================
# MainWindow
# =============================================================================

class MainWindow(QMainWindow):
    def __init__(self, name: str, version: str = "2.3.3", logger: Optional[logging.Logger] = None,
                 data_path: Union[PathLike[str], str, None] = None, *args, **kwargs):
        super().__init__()
        self.widget_name = name
        self.logger = logger
        self.data_path = Path.cwd()
        if data_path:
            self.data_path = Path(data_path)
        self.data_path.mkdir(exist_ok=True)
        if not self.logger:
            log_path = self.data_path / "script_log"
            log_path.mkdir(exist_ok=True)
            self.logger = create_logger(name=name, level=logging.DEBUG, log_path=log_path)
        self.logger.info(f"{name} logging started")
        if args:
            self.logger.debug(f"外部额外args参数：{args}")
        if kwargs:
            self.logger.debug(f"外部额外args参数：{kwargs}")

        self.setWindowTitle(f"BOM 工具 v{version}")
        self.resize(900, 600)

        if data_path:
            base = Path(data_path)
        else:
            base = Path.cwd()

        temp_dir = base / ".temp"
        temp_dir.mkdir(parents=True, exist_ok=True)

        output_dir = base / "result"
        output_dir.mkdir(parents=True, exist_ok=True)

        tabs = QTabWidget()
        self.setCentralWidget(tabs)

        self.integrate_tab = IntegrateTab(temp_dir, output_dir)
        tabs.addTab(self.integrate_tab, "Excel整合")

        self.match_tab = MatchTab(temp_dir, output_dir)
        tabs.addTab(self.match_tab, "Excel对比")


if __name__ == "__main__":
    app = QApplication(sys.argv)
    window = MainWindow(name="excel_match")
    window.show()
    sys.exit(app.exec())
